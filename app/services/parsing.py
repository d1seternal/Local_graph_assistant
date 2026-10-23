import multiprocessing
from pathlib import Path
from queue import Empty
import zipfile

from app.core.config import load_config


class ParsingError(ValueError):
    pass


def _extract(path: Path, extension: str, limits: dict) -> list[dict]:
    units = []
    characters = 0
    max_characters = int(limits.get("extracted_characters", 500_000))

    def add(text, source):
        nonlocal characters
        if not text.strip():
            return
        characters += len(text) + 1
        if characters > max_characters or len(units) >= 10_000:
            raise ParsingError("Превышен предел извлечённого текста.")
        units.append({"text": text, "source": source})

    if extension == ".txt":
        text = path.read_text(encoding="utf-8-sig")
        if len(text) > max_characters:
            raise ParsingError("Превышен предел извлечённого текста.")
        paragraph = []
        first_line = 1
        for number, line in enumerate(text.splitlines(), start=1):
            if line.strip():
                if not paragraph:
                    first_line = number
                paragraph.append(line)
            elif paragraph:
                add("\n".join(paragraph), {
                    "type": "txt", "line_start": first_line,
                    "line_end": number - 1,
                })
                paragraph = []
        if paragraph:
            add("\n".join(paragraph), {
                "type": "txt", "line_start": first_line,
                "line_end": first_line + len(paragraph) - 1,
            })

    elif extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ParsingError("PDF с шифрованием пока не поддерживается.")
        if len(reader.pages) > int(limits.get("pdf_pages", 50)):
            raise ParsingError("Превышен предел страниц PDF.")
        for number, page in enumerate(reader.pages, start=1):
            add(page.extract_text() or "", {"type": "pdf", "page": number})

    elif extension == ".docx":
        from docx import Document
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2000:
                raise ParsingError("Слишком много элементов внутри DOCX.")
            if sum(item.file_size for item in entries) > 50 * 1024 * 1024:
                raise ParsingError("Распакованный DOCX превышает 50 МиБ.")
            if "word/document.xml" not in archive.namelist():
                raise ParsingError("Файл не содержит основного текста DOCX.")
        document = Document(path)
        for number, paragraph in enumerate(document.paragraphs, start=1):
            add(paragraph.text, {"type": "docx_paragraph", "paragraph": number})
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                for column_number, cell in enumerate(row.cells, start=1):
                    add(cell.text, {
                        "type": "docx_cell", "table": table_number,
                        "row": row_number, "column": column_number,
                    })
    else:
        raise ParsingError("Формат не поддерживается.")

    if not units:
        raise ParsingError("Извлекаемый текст не найден. OCR пока не поддерживается.")
    return units


def _worker(path, extension, limits, output):
    try:
        output.put({"units": _extract(Path(path), extension, limits)})
    except ParsingError as error:
        output.put({"error": str(error)})
    except Exception as error:
        output.put({"error": "Не удалось разобрать документ: " + type(error).__name__})


def parse_document(path: Path, extension: str, timeout=30.0) -> list[dict]:
    context = multiprocessing.get_context("spawn")
    output = context.Queue(maxsize=1)
    process = context.Process(
        target=_worker,
        args=(str(path), extension, load_config().get("limits", {}), output),
        daemon=True,
    )
    process.start()
    try:
        try:
            result = output.get(timeout=timeout)
        except Empty as error:
            if process.is_alive():
                raise TimeoutError("Истекло время парсинга.") from error
            raise ParsingError("Процесс парсера завершился без результата.") from error
        if "error" in result:
            raise ParsingError(result["error"])
        return result["units"]
    finally:
        process.join(timeout=0.5)
        if process.is_alive():
            process.terminate()
            process.join(timeout=1)
        if process.is_alive():
            process.kill()
            process.join()
        output.close()
